from pathlib import Path
from docx import Document
from docx.shared import Pt

source = Path(r'C:\Users\kecac\Documents\UTN 2026\Cuatrimestre 3\Programacion 2 Jonathan moreno\wordLab\preguntasLab.docx')
doc = Document(source)
updates = {
    1: 'El objetivo es que el consumidor reciba cada dato del productor una sola vez, sin perder ni repetir valores.',
    4: 'No siempre. El productor puede reemplazar un dato antes de que se consuma, y el consumidor puede leer cuando el recipiente está vacío. Que funcione una vez no garantiza que siempre lo haga.',
    6: 'Los hilos no avanzan al mismo ritmo. Las pausas aleatorias cambian quién llega primero al recipiente, por eso una versión sin sincronización puede dar resultados distintos.',
    8: 'Ambos hilos deben usar el mismo recipiente. Se necesita synchronized para protegerlo, wait() para esperar y notifyAll() para avisar al otro hilo.',
    11: 'El productor espera si el recipiente está lleno y el consumidor espera si está vacío. La variable producido indica ese estado. Después de producir o consumir, se cambia la variable y se avisa al otro hilo. La espera va en un while para revisar la condición al despertar.',
    13: 'Porque 0 también es un dato válido. Puedo dejar amount en cero al consumir, pero producido es lo que indica si hay un dato disponible.',
    17: 'synchronized permite que solo un hilo a la vez entre a los métodos del recipiente. wait() pausa el hilo y libera el bloqueo. notify() despierta a uno de los que esperan; en mi código uso notifyAll(), que los despierta a todos para que revisen su condición.',
    19: 'En la última prueba no hubo excepciones. wait() y sleep() pueden lanzar InterruptedException si el hilo es interrumpido. También puede aparecer IllegalMonitorStateException si uso wait() o notify() sin tener el bloqueo del objeto.',
    21: 'Sí. Agregué synchronized y throws InterruptedException a ambos métodos. Como wait() puede lanzar esa excepción, las llamadas se manejan con try-catch en los hilos.',
    26: 'El programa compiló y, en la prueba, produjo y consumió del 0 al 9 en orden, sin pérdidas ni repeticiones. Ambos hilos terminaron. Falta contrastar los resultados con varias ejecuciones de la versión sin sincronización.',
    28: 'Quitar el if que depende de amount y esperar directamente con while (!producido). También escribir amount = 0 para que se entienda mejor y terminar el hilo si ocurre una interrupción.'
}
for i, text in updates.items():
    doc.paragraphs[i].text = text
for i in [24, 23, 22, 14]:
    p = doc.paragraphs[i]._element
    p.getparent().remove(p)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.0
doc.styles['Title'].font.size = Pt(19)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
doc.save(Path(__file__).parent / 'preguntasLab.docx')
print('Respuestas resumidas')
