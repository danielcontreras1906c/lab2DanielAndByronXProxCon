from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

out = Path(__file__).parent / 'preguntasLab.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.75)
sec.left_margin = sec.right_margin = Inches(0.85)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(21)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)

def question(text, answer):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.add_run(text).bold = True
    doc.add_paragraph(answer)

doc.add_paragraph('Laboratorio de productores y consumidores', 'Title')
doc.add_paragraph('Respuestas sobre la coordinación de dos hilos que comparten un recipiente de un solo dato. La sincronización debe permitir consumir cada valor producido exactamente una vez, aunque los hilos tengan pausas aleatorias diferentes.')
doc.add_paragraph('a Compilación y prueba inicial', 'Heading 1')
question('¿Funciona correctamente sin sincronización? ¿Qué problemas pueden aparecer?',
    'Sin sincronización no se garantiza un funcionamiento correcto. El productor puede depositar otro valor antes de que se retire el anterior y provocar una pérdida. El consumidor puede leer cuando no hay un dato disponible o repetir una lectura. Incluso si una ejecución parece correcta, el resultado depende del orden en que se ejecuten los hilos.')
question('¿Por qué pueden cambiar los resultados entre ejecuciones?',
    'Las pausas aleatorias de 1 a 2 segundos y la planificación de los hilos cambian el orden de las operaciones. Si el productor avanza más rápido, puede sobrescribir datos; si el consumidor avanza más rápido, puede intentar consumir un recipiente vacío. Estos son los resultados que se deben contrastar al ejecutar varias veces la versión sin sincronización.')
question('¿Qué mecanismos hacen falta para garantizar un funcionamiento correcto?',
    'Se necesita proteger el acceso al mismo recipiente y coordinar las condiciones de lleno y vacío. Ambos hilos deben compartir una única instancia de Container. Los métodos sincronizados protegen sus cambios; wait() permite esperar y notifyAll() avisa cuando el estado cambia. Thread.sleep() solamente introduce una pausa y no garantiza la coordinación.')
doc.add_paragraph('b Implementación de la sincronización', 'Heading 1')
question('¿Cómo se coordinan produce y consume?',
    'La variable producido indica si existe un dato disponible. El productor espera dentro de while (producido), guarda el número de la iteración y deja producido en true. El consumidor espera dentro de while (!producido), retira el dato y deja producido en false. Después de cada operación se notifica a los hilos que esperan. El while vuelve a comprobar la condición al despertar.')
question('¿Por qué distinguir el valor del estado del recipiente?',
    'El número 0 también puede ser producido. Por eso amount == 0 no basta para decidir si está vacío. Se puede dejar amount en cero después de consumir, pero producido debe controlar si hay un dato pendiente. El recipiente comienza vacío y produce asigna el dato recibido en lugar de acumularlo.')

doc.add_page_break()
doc.add_paragraph('c Preguntas de reflexión', 'Heading 1')
question('¿Qué métodos de sincronización se deben utilizar (wait(), notify(), synchronized)?',
    'synchronized es un modificador que hace que produce y consume utilicen el bloqueo del mismo recipiente: solo un hilo puede ejecutar esos métodos a la vez. wait() suspende al hilo y libera ese bloqueo mientras espera. notify() despierta a uno de los hilos que esperan; en el proyecto se utiliza notifyAll(), que los despierta a todos para que comprueben nuevamente su condición. El hilo despertado debe recuperar el bloqueo antes de continuar.')
question('¿El programa lanza alguna excepción? ¿Por qué?',
    'En la última prueba sincronizada no se observaron excepciones durante la ejecución. Sin embargo, wait() y Thread.sleep() pueden lanzar InterruptedException si el hilo es interrumpido. Al agregar throws InterruptedException a consume(), apareció un error de compilación porque su llamada estaba fuera del try; ese error se corrigió al capturar la excepción. También se produciría IllegalMonitorStateException si se llamara a wait(), notify() o notifyAll() sin poseer el bloqueo del objeto. En el proyecto se llaman dentro de métodos synchronized.')
question('¿Es necesario modificar el encabezado de los métodos produce y consume?',
    'Sí. En la implementación del proyecto ambos métodos deben ser synchronized y declarar throws InterruptedException, porque utilizan wait() y dejan que el hilo que los llama maneje la interrupción. Los encabezados actuales son los siguientes:')
for line in ['public synchronized void produce(int amount) throws InterruptedException',
             'public synchronized void consume() throws InterruptedException']:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(line)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
doc.add_paragraph('consume() puede devolver int si se quiere que Consumer reciba y muestre el valor retirado. En la versión actual devuelve void y el mensaje con el valor consumido se imprime dentro de Container.')
question('¿Qué se verificó al compilar y ejecutar nuevamente?',
    'La última versión revisada compiló y produjo los valores del 0 al 9. En la prueba se consumieron los mismos diez valores, en el mismo orden, sin pérdidas ni repeticiones; ambos hilos finalizaron. La consola mostró al consumidor esperando cuando no había un dato disponible. Una ejecución correcta no reemplaza la revisión de las condiciones ni las pruebas repetidas.')
question('¿Qué ajustes quedan para mejorar la versión revisada?',
    'Conviene quitar el if (amount <= 0) que rodea la espera del consumidor y usar directamente while (!producido). Para expresar que el recipiente queda en cero, amount = 0 es más claro que restar el atributo contra sí mismo. También conviene actualizar producido antes de notificar. Si ocurre una interrupción, se debe restaurar la marca con Thread.currentThread().interrupt() y terminar el hilo, evitando que el ciclo continúe y omita una operación.')

for root in (doc.styles.element, doc.element):
    for border in list(root.xpath('.//w:pBdr')):
        border.getparent().remove(border)
doc.save(out)
print(out)
